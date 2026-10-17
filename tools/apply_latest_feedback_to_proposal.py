"""Apply the 12 August 2026 supervisor-feedback revisions to the proposal.

This is a focused, repeatable edit of the current six-reference proposal. It
preserves the existing document structure and formatting while locking the
Part II SESOI, timeline, weather roles, Part III status, and fallback triggers.
"""

from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "Thesis_Proposal.docx"
OUTPUT = ROOT / "outputs" / "Thesis_Proposal.docx"


def find_paragraph(document, starts_with):
    prefixes = (
        starts_with if isinstance(starts_with, (tuple, list)) else (starts_with,)
    )
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith(tuple(prefixes))
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one paragraph starting with {prefixes!r}; "
            f"found {len(matches)}"
        )
    return matches[0]


def replace_paragraph_text(paragraph, text):
    source_runs = [run for run in paragraph.runs if run.text] or list(paragraph.runs)
    run_properties = (
        deepcopy(source_runs[0]._r.rPr)
        if source_runs and source_runs[0]._r.rPr is not None
        else None
    )
    for child in list(paragraph._p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def replace_cell_text(cell, text):
    paragraph = cell.paragraphs[0]
    replace_paragraph_text(paragraph, text)


def main():
    document = Document(SOURCE)

    replace_cell_text(document.tables[0].cell(4, 1), "August 2026")
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            for run in paragraph.runs:
                if "July 2026" in run.text:
                    run.text = run.text.replace("July 2026", "August 2026")

    abstract = find_paragraph(document, "This thesis addresses those questions")
    replace_paragraph_text(
        abstract,
        "This thesis addresses those questions through three connected analyses. "
        "Part I models summer Landsat land surface temperature using Sentinel-2 "
        "NDVI/vegetation greenness and built-form variables, with coarse "
        "soil-moisture information used only as a time-varying drought modifier. "
        "Part II uses recorded litres per 1,000 L/ha during the preceding 30 days "
        "as its primary exposure and proceeds confirmatorily only if the "
        "end-Month-2 simulation has at least 80% power for the pre-specified "
        "smallest effect of interest of -0.03 K per 1,000 L/ha; season-to-date "
        "dose and acute effects remain secondary. Part III tests whether logged "
        "watering intensity is distributed according to thermal exposure and "
        "socioeconomic vulnerability and produces a transparent vulnerability "
        "and watering-equity atlas at the finest defensible spatial scale.",
    )

    fallback = find_paragraph(document, "Part III is the guaranteed core deliverable")
    replace_paragraph_text(
        fallback,
        "Part III is the guaranteed core deliverable. The thesis will produce the "
        "watering-equity and vulnerability atlas even if Part II is underpowered "
        "or returns no detectable satellite-scale signal. Part II is conditional; "
        "if its end-Month-2 gate fails, its deliverable becomes the exposure audit, "
        "power result, and clearly labelled exploratory diagnostics. If the "
        "Landsat/Sentinel panel is not analysis-ready by the end of Month 2, Part I "
        "freezes at descriptive LST/greenness outputs and a data audit, and the "
        "Environmental Justice Atlas thermal layer supplies Part III.",
    )

    h3 = find_paragraph(document, "H3 — Equity.")
    replace_paragraph_text(
        h3,
        "H3 — Equity. Recorded care aligns more strongly with environmental need "
        "than with socioeconomic vulnerability, leaving some high-need areas "
        "underserved. The existing rho = +0.486, rho = -0.123, and ten-area "
        "screen are preliminary, descriptive, and hypothesis-generating because "
        "they informed H3. H3 becomes hypothesis-guided validation only after the "
        "care measures, transformations, weights, geography, exclusions, and "
        "multiplicity handling are locked; confirmatory is reserved for H1 and, "
        "if its gate passes, H2.",
    )

    h2 = find_paragraph(document, "H2 — Cumulative watering association.")
    replace_paragraph_text(
        h2,
        "H2 — Cumulative watering association. Among units with demonstrably "
        "high logging coverage, greater 30-day recorded litres per hectare are "
        "associated with lower LST, conditional on pre-period LST, greenness, "
        "built form, tree composition, weather, location, and acquisition. "
        "Season-to-date dose is secondary.",
    )

    weather = find_paragraph(document, "ESA CCI Soil Moisture Combined v09.2")
    replace_paragraph_text(
        weather,
        "ESA CCI Soil Moisture Combined v09.2 provides daily 0.25° values through "
        "2024 and is used only as a city-level or coarsely aggregated temporal "
        "drought modifier. ERA5 supplies the primary harmonised temporal controls: "
        "2 m temperature, dew point or derived humidity, total precipitation, "
        "10 m wind, surface pressure, and 7-, 14-, and 30-day antecedent summaries; "
        "ERA5 moisture and moisture-balance variables provide drought sensitivity. "
        "DWD station observations supply daily temperature, precipitation, "
        "humidity, wind, and sunshine-duration sensitivity and acquisition-day "
        "plausibility context. DWD is not the primary spatial weather control and "
        "does not directly validate satellite LST.",
    )

    gate = find_paragraph(
        document,
        ("Before outcome modelling, the Month 1 audit", "Month 1 audits exposure"),
    )
    replace_paragraph_text(
        gate,
        "Month 1 audits exposure, logging coverage, crosswalks, and usable-image "
        "counts. During Months 1–2, an outcome-only nuisance model estimates "
        "residual LST variance and spatial covariance without recorded-watering "
        "exposure, preferably using pre-platform summers or cross-fitting; the H2 "
        "coefficient remains masked until the gate. The primary exposure is "
        "recorded litres per 1,000 L/ha during the preceding 30 days at "
        "approximately 1 ha support. The SESOI is -0.03 K per additional 1,000 "
        "recorded L/ha/30 days, equivalent to about -0.10 K for 3,500 L/ha/30 "
        "days. This is a physical design scenario of roughly ten adequately "
        "watered tree-month equivalents, not an assumed or published effect. At "
        "two-sided alpha = 0.05, H2 proceeds confirmatorily only if end-Month-2 "
        "simulated power is at least 80% and the simulated MDE is no larger in "
        "magnitude than the SESOI. Otherwise Part II is exploratory.",
    )

    exposure = find_paragraph(
        document,
        ("The primary exposure is cumulative", "The primary exposure is the continuous"),
    )
    replace_paragraph_text(
        exposure,
        "The primary exposure is the continuous 30-day recorded dose per 1 ha "
        "imagery-support unit, scaled per 1,000 L/ha. Season-to-date dose, seven- "
        "and 14-day windows, acute effects, and logged-versus-no-record matching "
        "are secondary because no record does not establish no watering. "
        "Simulations and models preserve the observed continuous-dose distribution "
        "and spatial clustering. Models adjust for pre-period LST, contemporaneous "
        "and antecedent greenness, built form, tree composition, weather, "
        "acquisition, and location. Matching is retained only where overlap and "
        "covariate balance are acceptable. Spillover exclusions and shifted-date "
        "and shifted-location placebos remain primary credibility diagnostics.",
    )

    integration = find_paragraph(document, "The three parts use documented")
    replace_paragraph_text(
        integration,
        "The three parts use documented spatial crosswalks and scale-specific "
        "analytical units rather than forcing every dataset onto a common fine "
        "grid. Part I supplies median clear-sky summer LST and an adjusted "
        "hot-spot score to the Part III atlas. Part III displays those thermal "
        "inputs alongside low-greenness and socioeconomic components and reports "
        "care-gap sensitivity to weighting, geography, denominator, and thermal "
        "input. Part II does not determine the vulnerability index. Confirmatory "
        "testing is limited to H1 and conditionally H2; the existing Part III "
        "results are preliminary and H3 is hypothesis-guided validation after "
        "specification lock.",
    )

    objectives = document.tables[1]
    replace_cell_text(
        objectives.cell(2, 2),
        "Use 30-day recorded litres per hectare within high-coverage logged "
        "units as the primary exposure. Treat season-to-date dose, matched "
        "no-record comparisons, and acute effects as secondary diagnostics. "
        "Conditional on the end-Month-2 gate.",
    )

    mapping = document.tables[2]
    replace_cell_text(
        mapping.cell(1, 1),
        "Product-grid-cell/overpass panel model of Landsat LST with "
        "NDVI/greenness, built density, imperviousness, and water proximity. "
        "The 30 m grid is resampled product output, not native thermal detail. "
        "ERA5 provides primary temporal controls and drought sensitivity; DWD "
        "provides station sensitivity and plausibility context, not LST validation.",
    )
    replace_cell_text(
        mapping.cell(2, 1),
        "Use 30-day recorded litres per hectare within high-coverage logged units "
        "as the primary continuous dose; season-to-date dose is secondary. "
        "Confirmatory analysis proceeds only after the end-Month-2 gate. Matched "
        "no-record contrasts, 7- and 14-day windows, acute effects, and event-study "
        "diagnostics are secondary. Report balance, overlap, spillover exclusions, "
        "shifted-date/location placebos, and qualified adjusted associations.",
    )

    # Keep the data-source table on one page. The source and provider names remain
    # unambiguous while the shorter wording avoids a one-row spill page.
    sources = document.tables[3]
    replace_cell_text(sources.cell(13, 0), "Annual population and age structure")
    replace_cell_text(sources.cell(13, 1), "AfS Berlin-Brandenburg")
    replace_cell_text(
        sources.cell(13, 2),
        "Primary population, age, and denominator source.",
    )
    replace_cell_text(
        sources.cell(13, 3),
        "One annual LOR-level aggregate; published-geography limit.",
    )
    for row in sources.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_together = False
                for run in paragraph.runs:
                    run.font.size = Pt(6.6)

    plan = document.tables[4]
    replace_cell_text(
        plan.cell(1, 2),
        "Exposure/logging audit; source lock; crosswalks; start the blinded "
        "outcome-only Landsat nuisance panel; lock the Part III core.",
    )
    replace_cell_text(
        plan.cell(1, 3),
        "Protocol, exposure audit, and variance-estimation plan",
    )
    replace_cell_text(
        plan.cell(2, 2),
        "Complete the minimal Landsat panel and nuisance model without watering "
        "exposure; run the definitive gate at end Month 2. If the full panel is "
        "not ready, freeze descriptive thermal outputs and activate the EJA "
        "thermal fallback.",
    )
    replace_cell_text(
        plan.cell(2, 3),
        "End-Month-2 gate and Part I fallback decision",
    )
    replace_cell_text(
        plan.cell(4, 2),
        "If the gate passes and the Part III minimum is secure: run the 30-day "
        "continuous-dose model and credibility diagnostics. Otherwise report the "
        "exposure audit and exploratory diagnostics only.",
    )
    replace_cell_text(
        plan.cell(5, 2),
        "By end Month 3 protect the minimum core: Planungsraum plus district "
        "sensitivity, three care measures, MSS/EJA components, denominator "
        "sensitivity, and screening atlas. Defer extensions if behind.",
    )
    replace_cell_text(
        plan.cell(5, 3),
        "Guaranteed minimum core and screening atlas",
    )

    document.core_properties.modified = datetime.now(timezone.utc)
    document.core_properties.comments = (
        "Aligned with the most recent supervisor feedback dated August 2026; "
        "retains the five must-cites plus Egerer bibliography."
    )
    document.save(OUTPUT)
    if OUTPUT.stat().st_size == 0:
        raise RuntimeError(f"Output is empty: {OUTPUT}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
