"""Legacy entry point retained for compatibility.

The current proposal revision is implemented by
``apply_latest_feedback_to_proposal.py``. Running this file delegates to that
idempotent editor so the superseded July-only wording cannot be regenerated.
"""

from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs"
    / "Berlin_Urban_Heat_Thesis_Proposal_3-Part_Verified.docx"
)
OUTPUT = (
    ROOT
    / "outputs"
    / "Berlin_Urban_Heat_Thesis_Proposal_3-Part_Verified_Feedback-Incorporated.docx"
)


def find_paragraph(document, starts_with):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith(starts_with)
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one paragraph starting with {starts_with!r}; found {len(matches)}"
        )
    return matches[0]


def remove_text_runs(paragraph):
    for child in list(paragraph._p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            paragraph._p.remove(child)


def replace_paragraph_text(paragraph, text, bold_prefix=None):
    first_rpr = None
    last_rpr = None
    source_runs = [run for run in paragraph.runs if run.text]
    if not source_runs:
        source_runs = list(paragraph.runs)
    if source_runs:
        if source_runs[0]._r.rPr is not None:
            first_rpr = deepcopy(source_runs[0]._r.rPr)
        if source_runs[-1]._r.rPr is not None:
            last_rpr = deepcopy(source_runs[-1]._r.rPr)

    remove_text_runs(paragraph)
    if bold_prefix and text.startswith(bold_prefix):
        prefix_run = paragraph.add_run(bold_prefix)
        if first_rpr is not None:
            prefix_run._r.insert(0, deepcopy(first_rpr))
        prefix_run.bold = True

        body_run = paragraph.add_run(text[len(bold_prefix) :])
        if last_rpr is not None:
            body_run._r.insert(0, deepcopy(last_rpr))
        body_run.bold = False
    else:
        run = paragraph.add_run(text)
        if first_rpr is not None:
            run._r.insert(0, deepcopy(first_rpr))


def insert_paragraph_after(reference, text, style=None, bold_prefix=None):
    new_p = OxmlElement("w:p")
    if style is None and reference._p.pPr is not None:
        new_p.append(deepcopy(reference._p.pPr))
    reference._p.addnext(new_p)
    paragraph = Paragraph(new_p, reference._parent)
    if style is not None:
        paragraph.style = style
    replace_paragraph_text(paragraph, text, bold_prefix=bold_prefix)
    return paragraph


def copy_first_run_format(paragraph, reference):
    source_runs = [run for run in reference.runs if run.text]
    if not source_runs:
        source_runs = list(reference.runs)
    if not source_runs or source_runs[0]._r.rPr is None:
        return
    for run in paragraph.runs:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, deepcopy(source_runs[0]._r.rPr))


def replace_cell_text(cell, text):
    paragraph = cell.paragraphs[0]
    replace_paragraph_text(paragraph, text)
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)


def add_reference_after(reference_paragraph, text):
    return insert_paragraph_after(reference_paragraph, text)


def main():
    from apply_latest_feedback_to_proposal import main as apply_latest_feedback

    apply_latest_feedback()
    return

    document = Document(SOURCE)

    # Abstract: retain the existing three-part description while changing only
    # the Part II estimand, the decision gate, and the locked fallback.
    replace_paragraph_text(
        document.paragraphs[5],
        "This thesis addresses those questions through three connected analyses. "
        "Part I models summer Landsat land surface temperature using Sentinel-2 "
        "NDVI/vegetation greenness and built-form variables, with coarse "
        "soil-moisture information used only as a time-varying drought modifier. "
        "Part II treats 30-day and season-to-date logged litres as the primary "
        "exposure within high-logging-coverage units and proceeds to confirmatory "
        "modelling only if an up-front detectability analysis indicates at least "
        "80% simulated power; acute effects and matched no-record contrasts remain "
        "secondary diagnostics. Part III tests whether logged watering intensity "
        "is distributed according to thermal exposure and socioeconomic "
        "vulnerability and produces a transparent vulnerability and "
        "watering-equity atlas at the finest defensible spatial scale.",
    )
    replace_paragraph_text(
        document.paragraphs[6],
        "Land surface temperature is the primary outcome. DWD observations provide "
        "meteorological controls and broad plausibility checks, not direct "
        "validation of satellite LST. The watering analysis is observational: "
        "continuous cumulative dose within logged units is primary, while matching "
        "and event-study or difference-in-differences diagnostics are secondary and "
        "are used only where support and overlap are adequate. Results will be "
        "described as adjusted associations unless the full identification "
        "assumptions are supported. Part III is the guaranteed core deliverable; "
        "Part II is a conditional stretch analysis whose detectability result is "
        "itself reportable.",
    )

    # Make the data-science contribution explicit without changing the opening
    # narrative or adding a predictive-model claim.
    motivation_end = find_paragraph(document, "Yet the relevant chain of evidence")
    insert_paragraph_after(
        motivation_end,
        "Data-science positioning. The thesis deliberately prioritises multi-source "
        "geospatial data engineering, spatial inference, uncertainty "
        "quantification, and leakage-resistant validation over predictive-model "
        "novelty.",
        bold_prefix="Data-science positioning.",
    )

    # Reframe only the overstated second gap.
    gap_two = find_paragraph(document, "The drought interaction is often underspecified")
    replace_paragraph_text(
        gap_two,
        "Fine-scale specification of vegetation-water interactions. "
        "Drought-modulated vegetation cooling is an established literature, so the "
        "novelty is not the existence of the interaction. The contribution is a "
        "scale-aware Berlin analysis that separates fine spatial greenness from "
        "coarse temporal moisture and weather conditions (Schwaab et al., 2021; "
        "Gobatti et al., 2025).",
        bold_prefix="Fine-scale specification of vegetation-water interactions.",
    )

    gap_summary = find_paragraph(document, "The thesis brings these gaps together")
    fallback_heading = insert_paragraph_after(
        gap_summary,
        "2.1 Scope and Fallback Commitment",
        style="Heading 2",
    )
    insert_paragraph_after(
        fallback_heading,
        "Part III is the guaranteed core deliverable. The thesis will produce the "
        "multi-scale watering-equity and vulnerability atlas even if Part II is "
        "underpowered or returns no detectable satellite-scale signal. Part II is "
        "a conditional stretch analysis; if its detectability gate fails, its "
        "deliverable becomes the exposure audit, power result, and clearly labelled "
        "exploratory diagnostics. Part I remains the planned biophysical "
        "foundation.",
        style="Normal",
    )

    # Keep the existing objectives table and adjust only the Part II/III cells.
    objectives = document.tables[1]
    replace_cell_text(
        objectives.cell(2, 2),
        "Use cumulative 30-day and season-to-date dose within high-coverage logged "
        "units as the primary exposure. Treat matched no-record comparisons and "
        "acute effects as secondary diagnostics. Conditional on the detectability "
        "gate.",
    )
    replace_cell_text(
        objectives.cell(3, 2),
        "Compare catalogued-tree-normalised logged-watering rates with thermal and "
        "socioeconomic need using transparent component relationships, inequality "
        "measures, and geography sensitivity. Guaranteed core deliverable.",
    )

    # Research question, hypothesis, and analytical mapping for Part II.
    rq2 = find_paragraph(document, "RQ2.")
    replace_paragraph_text(
        rq2,
        "RQ2. Conditional on adequate detectability, is greater cumulative recorded "
        "watering intensity associated with lower local LST at satellite-supported "
        "units with demonstrably high logging coverage, after adjustment for "
        "pre-existing thermal, vegetation, built-form, weather, and location "
        "differences?",
    )
    h2 = find_paragraph(document, "H2 — Watering association.")
    replace_paragraph_text(
        h2,
        "H2 — Cumulative watering association. Among units with demonstrably high "
        "logging coverage, greater 30-day and season-to-date logged litres per "
        "satellite-supported unit are associated with lower LST, conditional on "
        "pre-period LST, greenness, built form, tree composition, weather, "
        "location, and acquisition.",
        bold_prefix="H2 — Cumulative watering association.",
    )
    mapping = document.tables[2]
    replace_cell_text(
        mapping.cell(2, 1),
        "Use cumulative 30-day and season-to-date logged litres within "
        "high-coverage logged units as the primary continuous dose. Confirmatory "
        "analysis proceeds only after a Month 1 power/detectability gate. Matched "
        "no-record contrasts, 7- and 14-day windows, acute effects, and event-study "
        "diagnostics are secondary. Report balance, overlap, spillover exclusions, "
        "shifted-date/location placebos, and qualified adjusted associations.",
    )

    contribution_end = find_paragraph(
        document, "A transparent spatial-justice assessment."
    )
    insert_paragraph_after(
        contribution_end,
        "McDonald et al. (2026) is used only as air-temperature and canopy-cover "
        "context; its result is not presented as evidence about Landsat LST.",
        style="Normal",
    )

    # Section 6: add the requested safeguards as local paragraphs around the
    # existing methods rather than replacing the section.
    units = find_paragraph(document, "The study covers all 12 administrative districts")
    replace_paragraph_text(
        units,
        units.text
        + " Part III repeats its main care-rate and need comparisons at "
        "Planungsraum, Bezirksregion, and district scales; this MAUP sensitivity is "
        "a headline robustness result rather than an appendix-only check.",
    )

    remote = find_paragraph(document, "Land surface temperature is obtained")
    insert_paragraph_after(
        remote,
        "For each Landsat date, Sentinel-2 NDVI is constructed from the nearest "
        "valid cloud-free scene or a short rolling composite centred on the "
        "overpass; window length and temporal distance are stored for audit. "
        "Because heat and water stress can also depress greenness, contemporaneous "
        "NDVI and LST are treated as jointly determined, and antecedent or lagged "
        "greenness provides a reverse-causality sensitivity check.",
    )

    drought = find_paragraph(document, "ESA CCI Soil Moisture Combined v09.2 provides")
    insert_paragraph_after(
        drought,
        "LST is observed only on clear-sky overpasses, selecting the sample toward "
        "hotter, drier, anticyclonic conditions. Acquisition and seasonal effects "
        "do not fully remove this mechanism. The analysis therefore reports the "
        "weather distribution of observed versus all summer days and checks "
        "coefficient stability across observed cloud fraction, antecedent "
        "cloudiness, air-mass, and humidity strata; inference is limited to the "
        "clear-sky acquisition regime.",
    )

    part_one = find_paragraph(document, "Spatial panel or mixed-effects models estimate")
    insert_paragraph_after(
        part_one,
        "Because greenness, imperviousness, and built density are near-substitutes, "
        "each enters a separate adjusted model before the full specification. "
        "Variance inflation, condition indices, and coefficient stability are "
        "reported. Model evaluation uses spatially blocked folds, with RMSE, MAE, "
        "and R-squared compared against random-fold results to quantify spatial "
        "leakage (Roberts et al., 2017).",
    )

    part_two_heading = find_paragraph(
        document, "6.6 Part II — Citizen Watering and Thermal Association"
    )
    detectability_heading = insert_paragraph_after(
        part_two_heading,
        "Up-Front Detectability Gate",
        style="Heading 3",
    )
    insert_paragraph_after(
        detectability_heading,
        "Before outcome modelling, the Month 1 audit uses actual usable-overpass "
        "counts, exposure density, crown-area proxies, residual LST variance, "
        "spatial correlation, and matching loss to estimate effective sample size, "
        "the minimum detectable association, and simulated power for 30-day and "
        "season-to-date dose. Confirmatory H2 proceeds only if the primary design "
        "has at least 80% simulated power for a practically plausible aggregated "
        "association. With adequate power, a near-zero estimate is substantive; "
        "without adequate power, the conclusion is non-detectability rather than "
        "evidence that watering has no physiological effect.",
        style="Normal",
    )

    part_two = find_paragraph(document, "Imagery units with recent recorded watering")
    replace_paragraph_text(
        part_two,
        "The primary exposure is cumulative logged watering intensity: litres in "
        "the 30 days before a valid overpass and season-to-date litres within an "
        "imagery-supported unit. The primary contrast is continuous dose within "
        "logged units and periods with demonstrably high platform coverage. "
        "Seven- and 14-day windows, acute effects, and logged-versus-no-record "
        "matching are secondary because no record does not establish no watering. "
        "Models adjust for pre-period LST, contemporaneous and antecedent "
        "greenness, built form, tree composition, weather, acquisition, and "
        "location. Matching is retained only where overlap and covariate balance "
        "are acceptable. A 150 m exclusion zone, with 100 m and 200 m sensitivity "
        "checks, reduces shared-microclimate and spillover risk. Shifted-date and "
        "shifted-location placebos remain primary credibility diagnostics; "
        "event-study or difference-in-differences plots remain diagnostic rather "
        "than causal proof.",
    )

    part_three = find_paragraph(document, "Watering is normalised using logged litres")
    insert_paragraph_after(
        part_three,
        "The analysis first shows each care metric's bivariate relationship with "
        "thermal need, vegetation deficit, socioeconomic status, older-population "
        "share, and population density so a composite index cannot hide trade-offs. "
        "The full analysis is repeated at Planungsraum, Bezirksregion, and district "
        "scales, and conclusions are labelled robust only when direction and "
        "substantive ranking are not artifacts of aggregation. Area-level "
        "relationships describe places, not individual volunteer motives. Because "
        "Baumkataster denominators cover only catalogued trees, the analysis maps "
        "denominator density, compares tree-, area-, and population-normalised "
        "rates, and reruns results after excluding weak-registry-coverage areas.",
    )

    integration = find_paragraph(
        document, "The three parts use documented spatial crosswalks"
    )
    replace_paragraph_text(
        integration,
        "The three parts use documented spatial crosswalks and scale-specific "
        "analytical units rather than forcing every dataset onto a common fine "
        "grid. Part I supplies median clear-sky summer LST and an adjusted residual "
        "hot-spot score to the Part III atlas. Part III displays those thermal "
        "inputs alongside low-greenness and socioeconomic components, then reports "
        "a transparent care-gap score with weighting, geography, denominator, and "
        "thermal-input sensitivity ranges. Part II does not determine the "
        "vulnerability index; if feasible, it adds a separate thermal-association "
        "evidence layer. Confirmatory tests are limited to the three primary "
        "hypotheses, while secondary coefficients are grouped by analysis family "
        "and controlled with the Benjamini-Hochberg false-discovery-rate procedure "
        "where formal multiplicity adjustment is warranted.",
    )

    # Preserve the six-row work plan while making the gate and fallback explicit.
    plan = document.tables[4]
    replace_cell_text(
        plan.cell(1, 2),
        "Literature review; source/version lock; audit Gieß den Kiez and "
        "Baumkataster coverage; define crosswalks; run the Part II exposure and "
        "detectability/power audit; lock the Part III core.",
    )
    replace_cell_text(plan.cell(1, 3), "Protocol, data audit, and Part II go/no-go memo")
    replace_cell_text(plan.cell(4, 0), "4 — Conditional Part II")
    replace_cell_text(
        plan.cell(4, 2),
        "If the gate passes: cumulative-dose models, overlap and balance checks, "
        "spillover exclusions, shifted-date/location placebos, and timing "
        "sensitivity. If not: exposure audit and exploratory diagnostics only.",
    )
    replace_cell_text(
        plan.cell(4, 3), "Qualified association evidence or detectability result"
    )
    replace_cell_text(plan.cell(5, 0), "5 — Core Part III")
    replace_cell_text(plan.cell(5, 1), "Months 2–5")
    replace_cell_text(plan.cell(5, 3), "Guaranteed core atlas")

    output_package = find_paragraph(
        document, "Thermal-association evidence package."
    )
    replace_paragraph_text(
        output_package,
        "Detectability and association evidence package. A pre-outcome power report "
        "and, only if justified, cumulative-dose models with overlap, balance, "
        "spillover, placebo, timing, and clear-sky diagnostics. A failed gate "
        "produces a documented non-detectability result rather than a failed "
        "thesis component.",
        bold_prefix="Detectability and association evidence package.",
    )

    limitations_end = find_paragraph(document, "Dated built-form inputs.")
    clear_sky_limit = insert_paragraph_after(
        limitations_end,
        "Clear-sky selection. Landsat samples hotter, drier, fair-weather "
        "conditions rather than all summer days. Cloudiness and air-mass "
        "sensitivity analyses describe this selection, and conclusions are limited "
        "to the observed acquisition regime.",
    )
    copy_first_run_format(clear_sky_limit, limitations_end)
    maup_limit = insert_paragraph_after(
        clear_sky_limit,
        "MAUP and ecological inference. Care rates, rankings, and need-care "
        "relationships may change with aggregation; area-level associations do not "
        "identify individual volunteer motives or behaviour.",
    )
    copy_first_run_format(maup_limit, limitations_end)

    # Add only the references introduced by the targeted feedback edits.
    references_end = find_paragraph(
        document,
        "Amt für Statistik Berlin-Brandenburg",
    )
    references_end = add_reference_after(
        references_end,
        "Gobatti, L., Bach, P. M., Maurer, M., and Leitão, J. P. (2025). "
        "Impact of soil moisture content on urban tree evaporative cooling and "
        "human thermal comfort. npj Urban Sustainability, 5, 26. "
        "https://doi.org/10.1038/s42949-025-00220-0.",
    )
    references_end = add_reference_after(
        references_end,
        "Openshaw, S. (1984). The Modifiable Areal Unit Problem. Concepts and "
        "Techniques in Modern Geography 38. Geo Books.",
    )
    references_end = add_reference_after(
        references_end,
        "Roberts, D. R. et al. (2017). Cross-validation strategies for data with "
        "temporal, spatial, hierarchical, or phylogenetic structure. Ecography, "
        "40, 913-929. https://doi.org/10.1111/ecog.02881.",
    )
    add_reference_after(
        references_end,
        "Schwaab, J. et al. (2021). The role of urban trees in reducing land "
        "surface temperatures in European cities. Nature Communications, 12, "
        "6763. https://doi.org/10.1038/s41467-021-26768-w.",
    )

    # The verified source has a manual page break immediately before Section 7.
    # The added safeguards otherwise leave Section 6.8 alone on a nearly empty
    # page, so remove only that one break and let the existing pagination flow.
    data_sources_heading = find_paragraph(document, "7. Data Sources")
    previous_element = data_sources_heading._p.getprevious()
    if previous_element is not None and previous_element.xpath(
        ".//w:br[@w:type='page']"
    ):
        previous_element.getparent().remove(previous_element)

    document.core_properties.modified = datetime.now(timezone.utc)
    document.core_properties.comments = (
        "Targeted revision incorporating supervisor feedback while preserving the "
        "verified proposal's structure and wording wherever possible."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    if OUTPUT.stat().st_size == 0:
        raise RuntimeError(f"Output is empty: {OUTPUT}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
