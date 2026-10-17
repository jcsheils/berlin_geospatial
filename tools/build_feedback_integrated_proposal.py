from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docx" / "Berlin_Urban_Heat_Thesis_Proposal_Feedback_Integrated.docx"

FONT = "Calibri"
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5C6670"
LIGHT = "F4F6F9"
LIGHTER = "FAFBFC"
BORDER = "C8D1DC"
WHITE = "FFFFFF"
BLACK = "111111"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_START_DXA = 120


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, size=None, bold=None, italic=None, color=BLACK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths):
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {TABLE_WIDTH_DXA}: {widths}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        keep_row_together(row)

    set_table_borders(table)


def format_cell_paragraph(paragraph, size=9, bold=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths, font_size=8.7, aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT)
        cell.text = header
        format_cell_paragraph(
            cell.paragraphs[0],
            size=9,
            bold=True,
            color=NAVY,
            align=(aligns[i] if aligns else WD_ALIGN_PARAGRAPH.LEFT),
        )
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = str(value)
            format_cell_paragraph(
                cells[i].paragraphs[0],
                size=font_size,
                color=BLACK,
                align=(aligns[i] if aligns else WD_ALIGN_PARAGRAPH.LEFT),
            )
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)
    return table


def add_page_field(paragraph):
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run2 = paragraph.add_run()
    set_run_font(run2, size=8.5, color=MUTED)
    run2._r.append(fld_char1)
    run2._r.append(instr_text)
    run2._r.append(fld_char2)


def setup_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = p.add_run("MASTER'S THESIS PROPOSAL")
    set_run_font(left, size=8.5, bold=True, color=MUTED)
    right = p.add_run("\tSaarland University | M.Sc. Data Science and Artificial Intelligence")
    set_run_font(right, size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left_footer = fp.add_run("FEEDBACK-INTEGRATED REVISION | AUGUST 2026")
    set_run_font(left_footer, size=8.5, color=MUTED)
    fp.add_run("\t")
    add_page_field(fp)


def set_style(style, size, color, bold, before, after, line_spacing=1.0, align=None):
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = rgb(color)
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing
    if align is not None:
        pf.alignment = align


def configure_styles(doc):
    normal = doc.styles["Normal"]
    set_style(normal, 11, BLACK, False, 0, 8, 1.333, WD_ALIGN_PARAGRAPH.JUSTIFY)
    normal.paragraph_format.widow_control = True

    set_style(doc.styles["Heading 1"], 16, BLUE, True, 18, 10, 1.0, WD_ALIGN_PARAGRAPH.LEFT)
    set_style(doc.styles["Heading 2"], 13, BLUE, True, 12, 6, 1.0, WD_ALIGN_PARAGRAPH.LEFT)
    set_style(doc.styles["Heading 3"], 12, DARK_BLUE, True, 8, 4, 1.0, WD_ALIGN_PARAGRAPH.LEFT)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles[name].paragraph_format.keep_together = True


def add_numbering(doc):
    numbering = doc.part.numbering_part.element

    def abstract_num(abstract_id, fmt, text, left, hanging, font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(left))
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "290")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), font)
            r_fonts.set(qn("w:hAnsi"), font)
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

    def num(num_id, abstract_id):
        node = OxmlElement("w:num")
        node.set(qn("w:numId"), str(num_id))
        abs_id = OxmlElement("w:abstractNumId")
        abs_id.set(qn("w:val"), str(abstract_id))
        node.append(abs_id)
        numbering.append(node)

    abstract_num(20, "bullet", "•", 540, 280, "Symbol")
    abstract_num(21, "decimal", "%1.", 540, 280)
    num(20, 20)
    num(21, 21)
    return 20, 21


def add_list_item(doc, text, num_id, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11)
    return p


def add_callout(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "20")
    left.set(qn("w:space"), "7")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    p_pr.append(borders)
    r1 = p.add_run(label + " ")
    set_run_font(r1, size=10.5, bold=True, color=NAVY)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=BLACK)
    return p


def add_body(doc, text, bold_lead=None, italic=False):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=11, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    for run in p.runs:
        set_run_font(
            run,
            size={1: 16, 2: 13, 3: 12}[level],
            bold=True,
            color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level],
        )
    return p


def add_title_page(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(8)
    run = kicker.add_run("MASTER'S THESIS PROPOSAL")
    set_run_font(run, size=10.5, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("Citizen Water, Vegetation Cooling, and Spatial Justice")
    set_run_font(run, size=23.5, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.keep_with_next = True
    run = subtitle.add_run("A Three-Part Spatiotemporal Analysis of Urban Heat Adaptation Across Berlin")
    set_run_font(run, size=13.5, color=MUTED)

    meta_rows = [
        ("Institution", "Saarland University - Department of Computer Science"),
        ("Program", "M.Sc. Data Science and Artificial Intelligence"),
        ("Supervisor", "Angela John, PhD Student, Interdisciplinary Institute for Societal Computing"),
        ("Candidate", "Joseph Carpenter Sheils"),
        ("Status", "Feedback-integrated revision"),
        ("Date", "August 2026"),
    ]
    table = doc.add_table(rows=0, cols=2)
    set_table_geometry(table, [1950, 7410])
    for label, value in meta_rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_shading(cells[0], LIGHT)
        format_cell_paragraph(cells[0].paragraphs[0], size=9.2, bold=True, color=NAVY)
        format_cell_paragraph(cells[1].paragraphs[0], size=9.2, color=BLACK)
    set_table_geometry(table, [1950, 7410])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("ABSTRACT")
    set_run_font(r, size=11, bold=True, color=BLUE)

    abstract = (
        "Berlin's heat burden is spatially uneven: dense, impervious neighbourhoods tend to experience "
        "higher land surface temperatures (LST), while vegetation provides cooling that can weaken when "
        "water is limited. Since 2020, the citizen platform Gieß den Kiez has recorded volunteer tree-watering "
        "events across the city. This thesis asks whether recorded care is distributed according to thermal "
        "and socioeconomic need and, only where the satellite record can support it, whether cumulative logged "
        "watering intensity has a detectable local LST association. Three linked analyses are proposed. Part I "
        "models summer Landsat LST using contemporaneous Sentinel-2 greenness, built form, and city-scale temporal "
        "drought modifiers. Part II is conditional on an end-Month-2 detectability assessment and prioritises "
        "30-day dose within high-logging-coverage units; the SESOI is -0.03 K per additional 1,000 recorded "
        "L/ha/30 days and acute event effects remain exploratory. Part III, "
        "the guaranteed core deliverable, evaluates watering equity at multiple spatial aggregations and produces "
        "a transparent vulnerability and care-gap atlas. The central data-science contribution is a reproducible, "
        "scale-aware integration and validation pipeline for heterogeneous geospatial data. The thesis does not "
        "equate LST with air temperature or felt heat and does not claim causal watering effects from incomplete "
        "observational logs."
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(abstract)
    set_run_font(r, size=10.2)
    p.add_run().add_break(WD_BREAK.PAGE)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    setup_header_footer(section)
    configure_styles(doc)
    bullet_id, decimal_id = add_numbering(doc)
    add_title_page(doc)

    add_heading(doc, "1. Motivation and Problem Statement")
    add_body(
        doc,
        "Berlin's heat exposure is unevenly distributed across its built environment. Impervious surfaces, "
        "dense urban form, limited vegetation, and population vulnerability combine to produce neighbourhoods "
        "with very different adaptation needs. Trees and other vegetation can reduce surface heating through "
        "shade and evapotranspiration, but the magnitude of that benefit depends on weather, water availability, "
        "urban form, and measurement scale (Schwaab et al., 2021; Gobatti et al., 2025)."
    )
    add_body(
        doc,
        "Berlin's Baumkataster contains more than 800,000 catalogued street- and park-tree records, although "
        "coverage of trees in green spaces is incomplete. Since May 2020, CityLAB's Gieß den Kiez platform has "
        "allowed residents to log watering dates, tree identities, locations, and litres. These records provide "
        "a rare trace of citizen-led climate adaptation, but they record platform activity rather than all water "
        "received by a tree. Municipal watering and unlogged resident care remain unobserved."
    )
    add_body(
        doc,
        "The proposal therefore separates three questions that require different units and standards of evidence: "
        "where summer surfaces are hottest and how greenness relates to LST; whether sufficiently dense cumulative "
        "logged watering has a detectable satellite-scale association; and whether recorded care is distributed "
        "relative to environmental and socioeconomic need. Districts are used for communication, while inference "
        "uses finer and explicitly documented spatial units."
    )
    add_callout(
        doc,
        "Data-science positioning.",
        "The thesis deliberately prioritises data engineering, spatial inference, uncertainty, and leakage-resistant "
        "validation over a predictive-model novelty claim. Its technical centrepiece is the reproducible integration "
        "of satellite, climate, tree-registry, citizen-generated, and socioeconomic data across incompatible spatial "
        "and temporal supports, with spatially blocked out-of-sample evaluation."
    )

    add_heading(doc, "2. Research Gap, Scope, and Fallback Commitment")
    add_heading(doc, "2.1 Research Gap", 2)
    add_list_item(
        doc,
        "Citizen watering as an environmental exposure. Resident watering and civic stewardship are established "
        "research topics (Moskell et al., 2016; Lachmund, 2022), and Egerer et al. (2024) identifies Gieß den Kiez "
        "as a Berlin coordination example. However, the logs have not been systematically aligned "
        "with valid satellite overpasses and defensible thermal support to test whether cumulative recorded care "
        "has a detectable LST association.",
        bullet_id,
        "Citizen watering as an environmental exposure.",
    )
    add_list_item(
        doc,
        "Fine-scale specification of vegetation-water interactions. Drought-modulated vegetation cooling is an "
        "established literature, so the novelty is not the existence of the interaction. The contribution is a "
        "scale-aware Berlin analysis that separates fine spatial greenness from coarse temporal moisture and weather "
        "conditions (Schwaab et al., 2021; Gobatti et al., 2025).",
        bullet_id,
        "Fine-scale specification of vegetation-water interactions.",
    )
    add_list_item(
        doc,
        "Spatial justice of adaptive care. Street-tree benefit inequalities are documented (Baró et al., 2019), but the distribution "
        "of recorded citizen care relative to thermal exposure and socioeconomic vulnerability is not.",
        bullet_id,
        "Spatial justice of adaptive care.",
    )

    add_heading(doc, "2.2 Locked Scope and Fallback Deliverable", 2)
    add_callout(
        doc,
        "Guaranteed core deliverable - Part III.",
        "The thesis will deliver the multi-scale watering-equity and vulnerability atlas even if Part II is "
        "underpowered or yields no detectable satellite-scale signal. The core can use Part I-derived thermal metrics "
        "when available and the Berlin Environmental Justice Atlas 2023/24 thermal-load and green-space indicators as "
        "a documented fallback. Part II is a conditional stretch analysis, not a condition for thesis completion."
    )
    add_body(
        doc,
        "Part I is a planned enabling analysis and substantive chapter. Part II proceeds as a confirmatory association "
        "study only if the end-Month-2 detectability gate is met. Otherwise, its deliverable is the detectability analysis, "
        "exposure audit, and clearly labelled exploratory results. This ordering protects the six-month schedule from "
        "the low probability that a transient, spatially diluted watering signal is recoverable in cloud-free Landsat observations."
    )

    add_heading(doc, "3. Research Objectives and Three-Part Design")
    add_table(
        doc,
        ["Part", "Primary purpose", "Planned approach and status"],
        [
            (
                "Part I\nUrban heat, vegetation, and drought",
                "Build the biophysical baseline and thermal-need layer.",
                "Model valid-overpass Landsat LST using contemporaneous Sentinel-2 greenness, built form, and temporal "
                "drought/weather modifiers. Planned enabling chapter.",
            ),
            (
                "Part II\nCitizen watering and thermal association",
                "Test whether cumulative logged care has a detectable imagery-scale association.",
                "Continuous cumulative dose within high-coverage logged units is primary. No-record comparisons, acute "
                "effects, and event-study diagnostics are secondary. Conditional stretch component.",
            ),
            (
                "Part III\nSpatial equity and vulnerability",
                "Measure whether recorded care reaches high-need places.",
                "Compare separate care-rate and need relationships, test geography sensitivity, and construct a transparent "
                "care-gap atlas. Guaranteed core deliverable.",
            ),
        ],
        [1800, 2500, 5060],
        font_size=8.8,
    )

    add_heading(doc, "4. Research Questions and Pre-Specified Hypotheses")
    add_heading(doc, "4.1 Research Questions", 2)
    add_list_item(
        doc,
        "RQ1. How do contemporaneous vegetation greenness, built density, imperviousness, proximity to water, and "
        "city-scale temporal drought conditions relate to summer LST across Berlin?",
        bullet_id,
    )
    add_list_item(
        doc,
        "RQ2. Conditional on adequate detectability, is greater cumulative recorded watering intensity associated "
        "with lower local LST at satellite-supported units, after adjustment for pre-existing thermal and built-form differences?",
        bullet_id,
    )
    add_list_item(
        doc,
        "RQ3. Is recorded citizen watering distributed in proportion to thermal and socioeconomic need, and which "
        "high-need areas show persistent care gaps across defensible spatial aggregations?",
        bullet_id,
    )

    add_heading(doc, "4.2 Primary Hypotheses", 2)
    add_list_item(
        doc,
        "H1 - Biophysical association. Higher contemporaneous greenness is associated with lower LST after adjustment "
        "for built form and acquisition conditions; this association is weaker during drier city-wide periods.",
        bullet_id,
        "H1 - Biophysical association.",
    )
    add_list_item(
        doc,
        "H2 - Cumulative watering association. Among units with demonstrably high logging coverage, greater 30-day "
        "and season-to-date logged litres per satellite-supported unit are associated with lower LST, conditional on "
        "pre-period LST, greenness, built form, tree composition, weather, location, and acquisition.",
        bullet_id,
        "H2 - Cumulative watering association.",
    )
    add_list_item(
        doc,
        "H3 - Equity. Recorded care aligns more strongly with environmental need than with socioeconomic vulnerability, "
        "leaving some high-need areas underserved. Existing Part III correlations and screened areas are preliminary, "
        "descriptive, and hypothesis-generating because they informed H3; validation begins only after the specification "
        "is locked.",
        bullet_id,
        "H3 - Equity.",
    )
    add_body(
        doc,
        "H1 and, if its gate passes, H2 are the confirmatory hypotheses. H3 is a hypothesis-guided validation analysis "
        "after care measures, transformations, weights, geography, exclusions, and multiplicity handling are locked. "
        "Secondary exposure windows, heterogeneity checks, placebos, and alternative index constructions are reported "
        "as robustness families."
    )

    add_heading(doc, "5. Analytical Mapping and Scientific Contribution")
    add_table(
        doc,
        ["RQ", "Primary estimand or output", "Method, validation, and decision rule"],
        [
            (
                "RQ1",
                "Adjusted greenness-LST association and greenness x temporal-dryness interaction.",
                "Product-grid-cell/overpass panel model; separate-predictor and full specifications; acquisition effects; "
                "Moran's I; spatially blocked cross-validation; clear-sky and NDVI-timing sensitivity.",
            ),
            (
                "RQ2",
                "Adjusted continuous association of cumulative logged litres with LST in high-coverage units.",
                "End-Month-2 power simulation at the locked SESOI; within-logged-unit continuous dose; matched no-record contrast secondary; "
                "100-200 m spillover exclusions; shifted-date and shifted-location placebos. Confirmatory only if the "
                "detectability gate passes.",
            ),
            (
                "RQ3",
                "Bivariate care-need slopes, inequality measures, and a transparent multi-scale care-gap score.",
                "Planungsraum, Bezirksregion, and district sensitivity; registry-coverage checks; ecological-inference "
                "caveat; equal-weight index and component-wise results shown together.",
            ),
        ],
        [900, 2380, 6080],
        font_size=8.5,
    )
    add_body(
        doc,
        "The contribution is methodological and substantive. Methodologically, the project develops a scale-aware "
        "spatiotemporal pipeline with reproducible crosswalks, exposure windows, blocked validation, and uncertainty "
        "audits. Substantively, it distinguishes the documented inequity of vegetation benefits from the less studied "
        "inequity of adaptive care (Baró et al., 2019)."
    )

    add_heading(doc, "6. Methodology")
    add_heading(doc, "6.1 Study Area, Units, and Spatial Support", 2)
    add_body(
        doc,
        "The study covers all 12 Berlin districts. Districts are presentation units, not the sole inferential sample. "
        "Part I operates at the Landsat product-grid-cell by overpass level while acknowledging that Landsat 8/9 thermal "
        "measurements have approximately 100 m native ground sampling and are resampled to a 30 m product grid. Part II "
        "uses approximately 100 m supported cells or justified local buffers and never interprets a product value as an "
        "individual-tree temperature. Part III uses LOR Planungsräume as the main geography where stable denominators "
        "exist, with Bezirksregion and district aggregation as headline aggregation-sensitivity analyses."
    )

    add_heading(doc, "6.2 Remote Sensing and Contemporaneous Greenness", 2)
    add_body(
        doc,
        "LST is obtained from Landsat 8 and 9 Collection 2 Level 2 surface-temperature products. Valid summer overpasses "
        "from 2013-2024 support descriptive mapping; the harmonised Sentinel-linked analytical panel covers 2017-2024. "
        "Sentinel-2 NDVI is constructed contemporaneously for each Landsat date using the nearest valid cloud-free scene "
        "or a short rolling composite centred on that date. Window length and temporal distance are stored for audit and "
        "tested in sensitivity analysis. NDVI is treated as vegetation greenness, not ready-made canopy cover."
    )
    add_body(
        doc,
        "Because heat and water stress can depress greenness, contemporaneous NDVI and LST may be jointly determined. "
        "The primary analysis is therefore associational. Antecedent or lagged greenness composites, where data support "
        "them, provide a reverse-causality sensitivity check. Dated urban-structure, imperviousness, land-use, and official "
        "waterbody layers are aligned through versioned spatial crosswalks."
    )

    add_heading(doc, "6.3 Drought, Weather, and Clear-Sky Selection", 2)
    add_body(
        doc,
        "ESA CCI Soil Moisture Combined v09.2 is a daily 0.25-degree product. Berlin spans only a few cells, so it is "
        "used as a city-wide or coarsely aggregated temporal modifier, never as a within-city moisture gradient. ERA5 "
        "supplies the primary harmonised temporal controls: 2 m temperature, dew point or derived humidity, total "
        "precipitation, 10 m wind, surface pressure, and 7-, 14-, and 30-day antecedent summaries; moisture and "
        "moisture-balance measures provide drought sensitivities. DWD station observations supply daily temperature, "
        "precipitation, humidity, wind, and sunshine-duration sensitivity and acquisition-day plausibility context. "
        "DWD is not the primary spatial weather control and does not directly validate LST."
    )
    add_body(
        doc,
        "LST is observed only on clear-sky overpasses. The observed sample is therefore selected toward hotter, drier, "
        "anticyclonic conditions. Acquisition and seasonal effects are necessary but not sufficient to remove this "
        "selection. The analysis will compare estimates across observed cloud fraction, antecedent cloudiness, air-mass "
        "and humidity strata; it will also report the weather distribution of observed versus all summer days. Inference "
        "is explicitly limited to the clear-sky acquisition regime."
    )

    add_heading(doc, "6.4 Part I - Urban Heat, Vegetation, and Drought", 2)
    add_body(
        doc,
        "Spatial panel or mixed-effects models estimate LST using greenness, built density, imperviousness, proximity to "
        "water, acquisition conditions, and temporal drought. Because greenness, imperviousness, and built density are "
        "near-substitutes, each is first entered in a separate adjusted model and then in a full model. Variance inflation, "
        "condition indices, coefficient stability, and partial effects are reported rather than hiding multicollinearity "
        "inside one specification."
    )
    add_body(
        doc,
        "Model evaluation uses spatially blocked folds that separate training and test areas because random folds can "
        "underestimate error when spatial dependence persists. Out-of-sample "
        "RMSE, MAE, and R-squared are reported for blocked folds and compared with random-fold results to quantify the "
        "optimism induced by spatial leakage. Moran's I, residual maps, and spatial lag/error alternatives diagnose "
        "remaining dependence."
    )

    add_heading(doc, "6.5 Up-Front Detectability Gate for Part II", 2)
    add_body(
        doc,
        "The primary Part II exposure is recorded litres per 1,000 L/ha during the preceding 30 days in a "
        "non-overlapping approximately 1 ha imagery-support unit. Month 1 audits exposure and logging coverage. During "
        "Months 1-2, an outcome-only nuisance model estimates residual LST variance and spatial covariance without "
        "recorded-watering exposure, preferably using pre-platform summers or cross-fitting; the H2 coefficient remains "
        "masked until the end-Month-2 gate. The project will calculate whether the observed watering density, usable clear-sky "
        "overpasses, thermal support, and residual LST variation permit a meaningful test. An illustrative dilution "
        "calculation shows the risk. A 100 m by 100 m footprint contains 10,000 m2. If one watered crown covers roughly "
        "30-80 m2, it occupies about 0.3-0.8% of the footprint. Even a local 1-3 °C surface difference would then average "
        "to only about 0.003-0.024 °C before point-spread, geolocation, weather, and mixed-surface noise. A cluster of ten "
        "similarly affected trees would imply roughly 0.03-0.24 °C. The locked SESOI is -0.03 K in support-mean LST "
        "per additional 1,000 recorded L/ha/30 days, equivalent to about -0.10 K for 3,500 L/ha/30 days. The latter is "
        "roughly ten adequately watered tree-month equivalents. These are planning scenarios, not assumed or published effects."
    )
    add_list_item(
        doc,
        "Estimate residual LST variance and spatial correlation from pre-exposure or low-exposure observations.",
        decimal_id,
    )
    add_list_item(
        doc,
        "Use actual valid-overpass counts, exposure densities, crown-area proxies, matching loss, and spatial clustering "
        "to simulate the effective sample size and minimum detectable association.",
        decimal_id,
    )
    add_list_item(
        doc,
        "At two-sided alpha = 0.05, calculate power for beta = -0.03 K per additional 1,000 recorded L/ha/30 days while "
        "preserving the observed continuous-dose distribution, clustering, and measurement-error scenarios. Treat "
        "season-to-date dose as secondary.",
        decimal_id,
    )
    add_list_item(
        doc,
        "Proceed with confirmatory H2 only if the primary design has at least 80% simulated power at the SESOI and the "
        "simulated MDE is no larger in magnitude. Otherwise report the detectability result and restrict Part II to exploratory "
        "description and method diagnostics.",
        decimal_id,
    )
    add_callout(
        doc,
        "Null-result commitment.",
        "With adequate power, a near-zero estimate and its confidence interval are a substantive finding about the absence "
        "of a detectable satellite-scale association. Without adequate power, the conclusion is limited to non-detectability; "
        "it is not evidence that watering has no physiological or microclimatic effect."
    )

    add_heading(doc, "6.6 Part II - Citizen Watering and Thermal Association", 2)
    add_body(
        doc,
        "The primary exposure is cumulative logged watering intensity: litres in the 30 days before a valid overpass and "
        "season-to-date litres within an imagery-supported unit. Seven- and 14-day windows are sensitivity specifications. "
        "Acute same-day or per-event effects are exploratory because watering effects are transient and exact within-day "
        "times are not consistently available. Same-day records are excluded unless a reliable timestamp proves that "
        "watering preceded acquisition."
    )
    add_body(
        doc,
        "The primary contrast is continuous dose within logged units and periods with demonstrably high platform coverage. "
        "A logging-coverage score based only on the exposure record - temporal continuity, active logging days, and the "
        "share of catalogued trees ever logged - is fixed before LST associations are examined. Logged-versus-no-record "
        "matching is secondary because no record does not mean no watering. Municipal crews and unlogged residents can "
        "contaminate comparison units and bias estimates toward zero."
    )
    add_body(
        doc,
        "Models adjust for pre-period LST, contemporaneous and antecedent greenness, built form, catalogued-tree composition, "
        "weather, acquisition, and location. Matching is retained only where overlap and covariate balance are acceptable. "
        "To reduce SUTVA and shared-microclimate violations, treated and comparison units are separated using a 150 m "
        "exclusion zone, with 100 m and 200 m sensitivity checks. Shifted-date and shifted-location placebos are primary "
        "credibility diagnostics. Event-study or difference-in-differences plots are reported only where repeated pre/post "
        "observations support them, and they remain diagnostic rather than causal proof."
    )

    add_heading(doc, "6.7 Part III - Spatial Equity and Vulnerability", 2)
    add_body(
        doc,
        "Recorded care is measured separately as logged litres per catalogued tree, the share of catalogued trees with a "
        "log, and logged events per 1,000 catalogued trees. The analysis first shows the bivariate relationships of each "
        "care metric with thermal need, vegetation deficit, socioeconomic status, older-population share, and population "
        "density. These plots and effect estimates remain visible alongside any composite index so that trade-offs are not hidden."
    )
    add_body(
        doc,
        "The main geography is the finest stable common LOR unit. The full analysis is repeated at Bezirksregion and "
        "district scales, and conclusions are labelled robust only if direction and substantive ranking are not artifacts "
        "of aggregation. This MAUP analysis is a headline result. Area-level associations describe places, not individual "
        "volunteer motives or behaviour; the ecological-inference limitation is stated in every interpretation."
    )
    add_body(
        doc,
        "Baumkataster denominators refer only to catalogued trees. Spatially varying catalogue completeness could bias "
        "care rates. The proposal therefore maps denominator density, documents green-space coverage gaps, compares tree-, "
        "area-, and population-normalised rates, and tests results after excluding areas with weak registry coverage. "
        "Monitoring Soziale Stadtentwicklung exclusions and exceptional land uses are handled through documented stable geographies."
    )

    add_heading(doc, "6.8 Quantitative Integration, Multiplicity, and Interpretation", 2)
    add_body(
        doc,
        "Part I produces two area-level thermal inputs for the atlas: median clear-sky summer LST exposure and an adjusted "
        "residual hot-spot score. Part III combines a standardised thermal score, low-greenness score, older-population share, "
        "and socioeconomic disadvantage into an equal-weight need index, while always displaying the components. A care "
        "score is derived from the three logged-watering rates after direction and scaling checks. The principal planning "
        "output is a care-gap score defined as standardised need minus standardised recorded care. Weighting, geography, "
        "denominator, and thermal-input alternatives are reported as sensitivity ranges."
    )
    add_body(
        doc,
        "Part II does not determine the vulnerability index. If feasible, it adds a separate evidence layer indicating "
        "whether cumulative recorded care has a detectable thermal association under the available satellite design. This "
        "avoids circularly defining need by an uncertain response model. Confirmatory tests are limited to H1 and, if its "
        "gate passes, H2. H3 is hypothesis-guided validation after specification lock. Secondary coefficients are grouped "
        "by analysis family, reported with effect sizes and confidence intervals, "
        "and controlled using the Benjamini-Hochberg false-discovery-rate procedure where formal multiplicity adjustment is warranted."
    )

    add_heading(doc, "7. Data Sources and Measurement Limits")
    add_body(
        doc,
        "Observation unit describes the delivered spatial or temporal unit of a value; it does not necessarily equal the "
        "instrument's native support. All datasets are version-locked in a manifest with licence, date range, geometry, "
        "coordinate reference system, missingness, and transformation metadata."
    )
    data_source_rows = [
        (
            "Landsat 8/9 C2 L2 ST\nUSGS / Google Earth Engine",
            "Primary LST outcome and thermal-need input.",
            "30 m product grid; approximately 100 m native thermal sampling; clear-sky only.",
            "I-III",
        ),
        (
            "Sentinel-2 surface reflectance\nESA Copernicus / GEE",
            "Contemporaneous NDVI and antecedent greenness sensitivity.",
            "10 m reflectance; composited around each Landsat date and summarised to outcome support.",
            "I-III",
        ),
        (
            "Land use, urban structure, imperviousness, waterbodies\nBerlin Senate / Geoportal",
            "Built-form controls and water proximity.",
            "Dated official layers; snapshots do not represent annual change.",
            "I-III",
        ),
        (
            "Berlin Baumkataster\nBerlin Geoportal",
            "Tree attributes and care-rate denominators.",
            "One catalogued tree; green-space and spatial coverage are incomplete.",
            "II-III",
        ),
        (
            "Gieß den Kiez logs\nBerlin Open Data / CityLAB",
            "Cumulative recorded-watering exposure and care metrics.",
            "One logged event; dates, litres, coordinates, tree ID; within-day time often unavailable; unlogged care absent.",
            "II-III",
        ),
        (
            "ESA CCI Soil Moisture Combined v09.2",
            "Primary temporal dryness modifier.",
            "Daily 0.25-degree cells; only a few cover Berlin; no neighbourhood gradient.",
            "I-II",
        ),
        (
            "ERA5\nCopernicus CDS",
            "Weather controls and alternative drought definitions.",
            "Coarse hourly grid aggregated to acquisition date or exposure window.",
            "I-II",
        ),
        (
            "DWD station observations",
            "Observed acquisition-day meteorology and plausibility context.",
            "Station-time values; not LST validation or tree-level temperature.",
            "I-II",
        ),
        (
            "Monitoring Soziale Stadtentwicklung 2025",
            "Socioeconomic vulnerability indicators.",
            "Published LOR values; small or exceptional areas may be excluded.",
            "III",
        ),
        (
            "Environmental Justice Atlas 2023/24",
            "Thermal-load and green-space benchmark; fallback thermal input.",
            "Categorical indicator per 2021 LOR Planungsraum.",
            "III",
        ),
        (
            "Einwohnerregister population and age",
            "Population, older-population share, and alternative denominators.",
            "Annual published aggregate by available LOR geography.",
            "III",
        ),
    ]
    data_table_headers = ["Dataset / provider", "Role", "Observation unit or principal limit", "Part"]
    data_table_widths = [2150, 2550, 3650, 1010]
    data_table_aligns = [
        WD_ALIGN_PARAGRAPH.LEFT,
        WD_ALIGN_PARAGRAPH.LEFT,
        WD_ALIGN_PARAGRAPH.LEFT,
        WD_ALIGN_PARAGRAPH.CENTER,
    ]
    add_table(
        doc,
        data_table_headers,
        data_source_rows[:4],
        data_table_widths,
        font_size=8.2,
        aligns=data_table_aligns,
    )
    add_table(
        doc,
        data_table_headers,
        data_source_rows[4:],
        data_table_widths,
        font_size=8.2,
        aligns=data_table_aligns,
    )

    add_heading(doc, "8. Work Plan, Decision Gates, and Feasibility")
    add_table(
        doc,
        ["Phase", "Period", "Tasks and decision gate", "Locked deliverable"],
        [
            (
                "1 - Protocol and coverage audit",
                "Month 1",
                "Version-lock data; audit exposure and registry coverage; define crosswalks; start the blinded outcome-only Landsat nuisance panel; lock the Part III care-rate foundation.",
                "Protocol, exposure audit, and variance-estimation plan",
            ),
            (
                "2 - Detectability gate",
                "Months 1-2",
                "Complete the nuisance model without watering exposure; simulate at the locked SESOI and issue the gate at end Month 2 before examining H2.",
                "End-Month-2 Part II go/no-go memo",
            ),
            (
                "3 - Part I",
                "Months 2-3",
                "Build the full contemporaneous panel; separate/full models; blocked CV; clear-sky and drought sensitivity. If not analysis-ready by end Month 2, freeze descriptive outputs and activate the EJA fallback.",
                "Thermal-need layer or documented fallback",
            ),
            (
                "4 - Core Part III",
                "Months 2-4",
                "Bivariate care-need analysis; MAUP and denominator sensitivity; vulnerability and care-gap atlas. At end Month 3 protect the minimum Planungsraum-plus-district core and defer extensions if behind.",
                "Guaranteed minimum core and screening atlas",
            ),
            (
                "5 - Conditional Part II",
                "Months 3-4",
                "If the gate passes and the Part III minimum is secure: cumulative dose models, balance, spillover exclusions, placebos, and timing sensitivity. Otherwise: exploratory audit only.",
                "Qualified association evidence or non-detectability result",
            ),
            (
                "6 - Synthesis and writing",
                "Months 5-6",
                "Integrate scores; multiplicity review; limitations appendix; policy interpretation; reproducibility release.",
                "Submitted thesis and code",
            ),
        ],
        [1550, 1050, 4820, 1940],
        font_size=8.3,
        aligns=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )
    add_callout(
        doc,
        "Schedule protection.",
        "At the end of Month 2, a missing analysis-ready Landsat/Sentinel panel freezes Part I at descriptive thermal "
        "outputs and activates the Environmental Justice Atlas thermal fallback. At the end of Month 3, protect the "
        "minimum Part III core before conditional Part II modelling; pause Part II even after a passed gate if that core "
        "is incomplete. A failed Part II gate reallocates time to Part III robustness, atlas quality, and synthesis."
    )

    add_heading(doc, "9. Expected Outputs")
    add_list_item(
        doc,
        "Master's thesis. A three-part structure with Part III guaranteed, Part I planned, and Part II clearly conditional; "
        "the results chapters are supported by a limitations and diagnostics appendix.",
        bullet_id,
        "Master's thesis.",
    )
    add_list_item(
        doc,
        "Berlin vulnerability and watering-equity atlas. Multi-scale maps, separate care-need relationships, care-gap scores, "
        "denominator warnings, and a transparent indicator dictionary.",
        bullet_id,
        "Berlin vulnerability and watering-equity atlas.",
    )
    add_list_item(
        doc,
        "Detectability and association evidence package. A pre-outcome power report and, if justified, cumulative-dose models "
        "with balance, blocked validation, spillover, placebo, timing, and clear-sky diagnostics.",
        bullet_id,
        "Detectability and association evidence package.",
    )
    add_list_item(
        doc,
        "Reproducible code and data documentation. Versioned scripts, dataset manifest, spatial crosswalks, tests, model cards, "
        "and documented licences for all releasable products.",
        bullet_id,
        "Reproducible code and data documentation.",
    )

    add_heading(doc, "10. Software and Technical Environment")
    add_table(
        doc,
        ["Tool or library", "Purpose"],
        [
            ("Google Earth Engine (Python API)", "Landsat and Sentinel retrieval, QA masking, compositing, and export."),
            ("Python: GeoPandas, Rasterio, Pandas, NumPy", "Spatial integration, crosswalks, exposure construction, and quality assurance."),
            ("Python: statsmodels, linearmodels, scikit-learn", "Panel models, diagnostics, blocked validation, and simulation."),
            ("R: fixest, lme4, MatchIt, spatialreg", "Association models, balance, event-study diagnostics, and spatial sensitivity."),
            ("PySAL / GeoDa", "Spatial weights, Moran's I, local clustering, and spatial regression."),
            ("QGIS", "Visual geometry checks and final atlas cartography."),
            ("Git / GitHub", "Version control, tests, environment capture, and reproducibility."),
        ],
        [2700, 6660],
        font_size=8.8,
    )

    add_heading(doc, "11. Ethics, Interpretation, and Limitations")
    add_heading(doc, "11.1 Data Ethics", 2)
    add_body(
        doc,
        "The public Gieß den Kiez data record tree-level events rather than volunteer identities. Socioeconomic and "
        "population measures are used only in published aggregate form. Maps and tables will be checked for small-cell "
        "disclosure risk, and licence and attribution requirements will be documented. Formal university ethics requirements "
        "will be confirmed before analysis."
    )
    add_heading(doc, "11.2 Interpretation Boundaries", 2)
    limitations = [
        (
            "LST is not air temperature or felt heat.",
            "The study estimates surface-temperature relationships during clear-sky satellite acquisitions. DWD observations "
            "provide meteorological context, not direct validation.",
        ),
        (
            "Recorded watering is incomplete and non-random.",
            "Platform logs omit municipal and unlogged resident care. The absence of a record is never labelled unwatered.",
        ),
        (
            "Acute effects are poorly timed for Landsat.",
            "Revisit intervals and cloud masking make short post-event effects unlikely to be recoverable. Cumulative dose is primary.",
        ),
        (
            "Drought has little within-city spatial leverage.",
            "ESA CCI and ERA5 moisture variables are temporal modifiers and are collinear with heat and season; alternative definitions "
            "and acquisition controls are required.",
        ),
        (
            "Clear-sky selection limits generalisation.",
            "Results apply to observed clear-sky conditions. Cloudiness and air-mass sensitivity analyses describe this selection.",
        ),
        (
            "MAUP and ecological inference constrain equity claims.",
            "Rates and rankings may change with aggregation, and area-level relationships do not identify individual motivations.",
        ),
        (
            "Tree-registry completeness may bias denominators.",
            "Catalogue coverage is mapped and alternative denominators and restricted samples are reported.",
        ),
    ]
    for lead, rest in limitations:
        add_list_item(doc, f"{lead} {rest}", bullet_id, lead)

    add_heading(doc, "12. Key References")
    references = [
        "Baró, F., Calderón-Argelich, A., Langemeyer, J., and Connolly, J. J. T. (2019). Under one canopy? Assessing the distributional environmental justice implications of street tree benefits in Barcelona. Environmental Science & Policy, 102, 54-64. https://doi.org/10.1016/j.envsci.2019.08.016.",
        "Egerer, M., Schmack, J. M., Vega, K., Ordóñez Barona, C., and Raum, S. (2024). The challenges of urban street trees and how to overcome them. Frontiers in Sustainable Cities, 6, 1394056. https://doi.org/10.3389/frsc.2024.1394056.",
        "Gobatti, L., Bach, P. M., Maurer, M., and Leitão, J. P. (2025). Impact of soil moisture content on urban tree evaporative cooling and human thermal comfort. npj Urban Sustainability, 5, 26. https://doi.org/10.1038/s42949-025-00220-0.",
        "Lachmund, J. (2022). Stewardship practice and the performance of citizenship: Greening tree-pits in the streets of Berlin. Environment and Planning C: Politics and Space, 40(6), 1290-1306. https://doi.org/10.1177/23996544211070204.",
        "Moskell, C., Bassuk, N., Allred, S., and MacRae, P. (2016). Engaging residents in street tree stewardship: Results of a tree watering outreach intervention. Arboriculture & Urban Forestry, 42(5), 301-317. https://doi.org/10.48044/jauf.2016.027.",
        "Schwaab, J., Meier, R., Mussetti, G., Seneviratne, S. I., Bürgi, C., and Davin, E. L. (2021). The role of urban trees in reducing land surface temperatures in European cities. Nature Communications, 12, 6763. https://doi.org/10.1038/s41467-021-26768-w.",
    ]
    for ref_index, ref in enumerate(references):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(ref)
        set_run_font(r, size=9.2)

    add_heading(doc, "Approval and Signatures")
    add_table(
        doc,
        ["Submitted by (Candidate)", "Approved by (Supervisor)"],
        [
            (
                "\n______________________________\nJoseph Carpenter Sheils\nDate: ___________________",
                "\n______________________________\nAngela John, PhD Student\nDate: ___________________",
            )
        ],
        [4680, 4680],
        font_size=9.2,
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    props = doc.core_properties
    props.title = "Citizen Water, Vegetation Cooling, and Spatial Justice"
    props.subject = "Feedback-integrated master's thesis proposal"
    props.author = "Joseph Carpenter Sheils"
    props.keywords = "Berlin, urban heat, citizen watering, spatial justice, land surface temperature"
    props.comments = "Revised to incorporate supervisor feedback dated August 2026."

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
