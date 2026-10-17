from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "Thesis_Proposal.docx"
OUTPUT = ROOT / "outputs" / "Thesis_Proposal.docx"


REFERENCES = [
    (
        "Baró, F., Calderón-Argelich, A., Langemeyer, J., and Connolly, "
        "J. J. T. (2019). Under one canopy? Assessing the distributional "
        "environmental justice implications of street tree benefits in "
        "Barcelona. Environmental Science & Policy, 102, 54–64. "
        "https://doi.org/10.1016/j.envsci.2019.08.016."
    ),
    (
        "Egerer, M., Schmack, J. M., Vega, K., Ordóñez Barona, C., and Raum, "
        "S. (2024). The challenges of urban street trees and how to overcome "
        "them. Frontiers in Sustainable Cities, 6, 1394056. "
        "https://doi.org/10.3389/frsc.2024.1394056."
    ),
    (
        "Gobatti, L., Bach, P. M., Maurer, M., and Leitão, J. P. (2025). "
        "Impact of soil moisture content on urban tree evaporative cooling "
        "and human thermal comfort. npj Urban Sustainability, 5, 26. "
        "https://doi.org/10.1038/s42949-025-00220-0."
    ),
    (
        "Lachmund, J. (2022). Stewardship practice and the performance of "
        "citizenship: Greening tree-pits in the streets of Berlin. Environment "
        "and Planning C: Politics and Space, 40(6), 1290–1306. "
        "https://doi.org/10.1177/23996544211070204."
    ),
    (
        "Moskell, C., Bassuk, N., Allred, S., and MacRae, P. (2016). Engaging "
        "residents in street tree stewardship: Results of a tree watering "
        "outreach intervention. Arboriculture & Urban Forestry, 42(5), "
        "301–317. https://doi.org/10.48044/jauf.2016.027."
    ),
    (
        "Schwaab, J., Meier, R., Mussetti, G., Seneviratne, S. I., Bürgi, C., "
        "and Davin, E. L. (2021). The role of urban trees in reducing land "
        "surface temperatures in European cities. Nature Communications, 12, "
        "6763. https://doi.org/10.1038/s41467-021-26768-w."
    ),
]


def find_paragraph(document, starts_with):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith(starts_with)
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one paragraph starting with {starts_with!r}; "
            f"found {len(matches)}"
        )
    return matches[0]


def find_optional_paragraph(document, starts_with):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith(starts_with)
    ]
    if len(matches) > 1:
        raise ValueError(
            f"Expected at most one paragraph starting with {starts_with!r}; "
            f"found {len(matches)}"
        )
    return matches[0] if matches else None


def remove_text_runs(paragraph):
    for child in list(paragraph._p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            paragraph._p.remove(child)


def replace_paragraph_text(paragraph, text, bold_prefix=None):
    source_runs = [run for run in paragraph.runs if run.text]
    if not source_runs:
        source_runs = list(paragraph.runs)
    first_rpr = (
        deepcopy(source_runs[0]._r.rPr)
        if source_runs and source_runs[0]._r.rPr is not None
        else None
    )
    last_rpr = (
        deepcopy(source_runs[-1]._r.rPr)
        if source_runs and source_runs[-1]._r.rPr is not None
        else None
    )

    remove_text_runs(paragraph)
    if bold_prefix and text.startswith(bold_prefix):
        prefix_run = paragraph.add_run(bold_prefix)
        if first_rpr is not None:
            prefix_run._r.insert(0, deepcopy(first_rpr))
        prefix_run.bold = True

        body_run = paragraph.add_run(text[len(bold_prefix) :])
        if last_rpr is not None:
            body_run._r.insert(0, deepcopy(last_rpr))
        body_run.bold = False
    else:
        run = paragraph.add_run(text)
        if first_rpr is not None:
            run._r.insert(0, first_rpr)


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def insert_formatted_paragraph_before(reference, template, text):
    new_p = OxmlElement("w:p")
    if template._p.pPr is not None:
        new_p.append(deepcopy(template._p.pPr))
    reference._p.addprevious(new_p)
    paragraph = Paragraph(new_p, reference._parent)
    paragraph.add_run(text)

    template_runs = [run for run in template.runs if run.text]
    if template_runs and template_runs[0]._r.rPr is not None:
        paragraph.runs[0]._r.insert(0, deepcopy(template_runs[0]._r.rPr))
    return paragraph


def replace_reference_section(document):
    heading = find_paragraph(document, "12. Key References")
    approval = find_paragraph(document, "Approval and Signatures")
    paragraphs = list(document.paragraphs)
    heading_index = next(
        index for index, paragraph in enumerate(paragraphs) if paragraph._p is heading._p
    )
    approval_index = next(
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph._p is approval._p
    )
    old_block = paragraphs[heading_index + 1 : approval_index]
    template = next(paragraph for paragraph in old_block if paragraph.text.strip())

    for paragraph in old_block:
        remove_paragraph(paragraph)
    for reference in REFERENCES:
        insert_formatted_paragraph_before(approval, template, reference)


def validate_revision(document):
    paragraphs = list(document.paragraphs)
    heading = find_paragraph(document, "12. Key References")
    approval = find_paragraph(document, "Approval and Signatures")
    heading_index = next(
        index for index, paragraph in enumerate(paragraphs) if paragraph._p is heading._p
    )
    approval_index = next(
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph._p is approval._p
    )
    block = paragraphs[heading_index + 1 : approval_index]
    entries = [paragraph.text.strip() for paragraph in block if paragraph.text.strip()]
    if entries != REFERENCES:
        raise RuntimeError("Reference section does not match the six-item shortlist.")

    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    required_citations = (
        "Schwaab et al., 2021",
        "Gobatti et al., 2025",
        "Moskell et al., 2016",
        "Baró et al., 2019",
        "Lachmund, 2022",
        "Egerer et al. (2024)",
    )
    missing = [citation for citation in required_citations if citation not in full_text]
    if missing:
        raise RuntimeError("Missing in-text citations: " + ", ".join(missing))
    for orphan in ("McDonald et al. (2026)", "Roberts et al., 2017"):
        if orphan in full_text:
            raise RuntimeError(f"Removed reference remains cited: {orphan}")


def main():
    document = Document(SOURCE)

    citizen_gap = find_paragraph(
        document, "Citizen watering is rarely modelled as an environmental exposure."
    )
    replace_paragraph_text(
        citizen_gap,
        "Citizen watering is rarely modelled as an environmental exposure. "
        "Resident tree watering and stewardship have been studied through "
        "outreach and civic-care lenses (Moskell et al., 2016; Lachmund, 2022), "
        "and Egerer et al. (2024) identify Gieß den Kiez as a Berlin "
        "coordination example. However, the platform record has not been "
        "systematically aligned with satellite overpasses to test whether recent "
        "recorded watering is associated with a detectable local LST difference.",
        bold_prefix="Citizen watering is rarely modelled as an environmental exposure.",
    )

    justice_gap = find_paragraph(
        document, "The spatial justice of adaptive care is not well quantified."
    )
    replace_paragraph_text(
        justice_gap,
        "The spatial justice of adaptive care is not well quantified. "
        "Distributional inequalities in street-tree benefits are documented "
        "(Baró et al., 2019), but it remains unclear whether recorded watering "
        "follows environmental need, socioeconomic capacity, or both at "
        "neighbourhood scale.",
        bold_prefix="The spatial justice of adaptive care is not well quantified.",
    )

    orphan = find_optional_paragraph(document, "McDonald et al. (2026)")
    if orphan is not None:
        remove_paragraph(orphan)

    validation = find_paragraph(document, "Because greenness, imperviousness")
    replace_paragraph_text(
        validation,
        "Because greenness, imperviousness, and built density are near-substitutes, "
        "each enters a separate adjusted model before the full specification. "
        "Variance inflation, condition indices, and coefficient stability are "
        "reported. Model evaluation will use spatially blocked folds, with RMSE, "
        "MAE, and R-squared compared against random-fold results as a diagnostic "
        "for spatial leakage.",
    )

    replace_reference_section(document)
    validate_revision(document)

    document.core_properties.modified = datetime.now(timezone.utc)
    document.core_properties.comments = (
        "Interim bibliography revision retaining the five substantive must-cites "
        "plus Egerer et al. (2024)."
    )
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    if OUTPUT.stat().st_size == 0:
        raise RuntimeError(f"Output is empty: {OUTPUT}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
